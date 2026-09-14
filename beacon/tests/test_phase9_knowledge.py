"""BEACON V2 — Phase 9 Company Knowledge Retrieval / RAG Infrastructure Test Suite.

Verifies end-to-end knowledge lifecycle: upload, ingestion, structural extraction,
deterministic chunking, SQLite persistence, scoped BM25 retrieval, provenance tracking,
and REST API endpoints.
"""

import io
from pathlib import Path
import sys
import pytest
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

import docx
from api.main import app
from auth import UserRole, auth_repository, create_access_token
from knowledge.models import (
    DocumentMetadata,
    DocumentStatus,
    DocumentType,
    KnowledgeContext,
    RetrievalPolicy,
    RetrievalStatus,
)
from knowledge.persistence import KnowledgePersistenceStore, knowledge_persistence
from knowledge.service import KnowledgeService, knowledge_service


def get_admin_headers() -> dict[str, str]:
    """Ensures an active system admin exists and returns valid authorization headers."""
    admin = auth_repository.get_user_by_username("admin")
    if not admin:
        admin = auth_repository.create_user(
            username="admin",
            password="AdminPassword2026!",
            is_system_admin=True,
        )
    token = create_access_token(user=admin)
    return {"Authorization": f"Bearer {token}"}


def make_test_pdf(text: str = "Enterprise Reimbursement Policy. Section 1. Travel expenses require itemized receipts.") -> bytes:
    """Generates valid synthetic PDF binary bytes with embedded text."""
    pdf_str = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length {len(text) + 60} >>
stream
BT
/F1 12 Tf
72 712 Td
({text}) Tj
ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000414 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
502
%%EOF"""
    return pdf_str.encode("latin-1")


def make_test_docx(title: str = "Retail Discount Rules", paragraphs: list[str] = None) -> bytes:
    """Generates valid synthetic DOCX binary bytes with paragraphs and headings."""
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for p in (paragraphs or ["VIP Gold tier customers receive 20% discount on electronics."]):
        doc.add_paragraph(p)
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


class TestPhase9CompanyKnowledge:
    """Comprehensive test suite for Phase 9 Knowledge Ingestion & RAG Retrieval."""

    def test_pdf_document_ingestion_and_chunking(self, tmp_path):
        """Valid PDF is parsed, normalized, chunked, and stored in SQLite as READY."""
        store = KnowledgePersistenceStore(db_path=str(tmp_path / "test_pdf.sqlite3"))
        service = KnowledgeService(storage_root=tmp_path, persistence=store)
        pdf_bytes = make_test_pdf("Company Travel Guidelines. Hotel bookings must not exceed 250 dollars per night.")

        doc_meta = service.ingest_document(
            content=pdf_bytes,
            filename="travel_policy.pdf",
            application_id="app_finance_01",
            chunk_size=200,
            chunk_overlap=50,
        )

        assert doc_meta.status == DocumentStatus.READY
        assert doc_meta.document_type == DocumentType.PDF
        assert doc_meta.filename == "travel_policy.pdf"
        assert doc_meta.page_count >= 1
        assert doc_meta.chunk_count >= 1

        # Check persisted chunks
        chunks = service.list_chunks(doc_meta.document_id)
        assert len(chunks) == doc_meta.chunk_count
        assert any("Hotel bookings" in c.text for c in chunks)
        assert chunks[0].metadata.document_id == doc_meta.document_id
        assert chunks[0].metadata.application_id == "app_finance_01"

    def test_docx_document_ingestion_and_chunking(self, tmp_path):
        """Valid DOCX with heading and paragraph is ingested, chunked, and indexed."""
        store = KnowledgePersistenceStore(db_path=str(tmp_path / "test_docx.sqlite3"))
        service = KnowledgeService(storage_root=tmp_path, persistence=store)
        docx_bytes = make_test_docx(
            title="Healthcare Benefit Guide",
            paragraphs=["Preventative care services are covered at 100 percent with zero copay."],
        )

        doc_meta = service.ingest_document(
            content=docx_bytes,
            filename="benefits_guide.docx",
            application_id="app_hr_01",
            chunk_size=250,
            chunk_overlap=50,
        )

        assert doc_meta.status == DocumentStatus.READY
        assert doc_meta.document_type == DocumentType.DOCX
        assert doc_meta.chunk_count >= 1

        chunks = service.list_chunks(doc_meta.document_id)
        assert len(chunks) > 0
        assert any("Preventative care" in c.text for c in chunks)

    def test_scoped_knowledge_retrieval(self, tmp_path):
        """Retrieval strictly searches only the target application scope and returns grounded context."""
        store = KnowledgePersistenceStore(db_path=str(tmp_path / "test_ret.sqlite3"))
        service = KnowledgeService(storage_root=tmp_path, persistence=store)
        pdf_bytes = make_test_pdf("Retail pricing rules. Tier 1 loyalty members receive free expedited shipping.")

        service.ingest_document(
            content=pdf_bytes,
            filename="retail_shipping.pdf",
            application_id="app_retail_01",
        )

        # 1. Matching query in correct scope -> SUCCESS_WITH_RESULTS
        ctx = service.retrieve(
            query="loyalty members shipping",
            application_id="app_retail_01",
            top_k=3,
        )
        assert ctx.status == RetrievalStatus.SUCCESS_WITH_RESULTS
        assert len(ctx.retrieved_chunks) >= 1
        assert "retail_shipping.pdf" in ctx.sources
        assert ctx.retrieved_chunks[0].relevance_score > 0.0
        assert ctx.retrieval_fingerprint != ""

        # 2. Non-existent scope -> SUCCESS_NO_RESULTS (Cross-scope isolation)
        ctx_other = service.retrieve(
            query="loyalty members shipping",
            application_id="app_unrelated_99",
            top_k=3,
        )
        assert ctx_other.status == RetrievalStatus.SUCCESS_NO_RESULTS
        assert len(ctx_other.retrieved_chunks) == 0

    def test_duplicate_document_content_detection(self, tmp_path):
        """Ingesting exact same document content hash in same scope flags DUPLICATE and avoids duplicate chunking."""
        store = KnowledgePersistenceStore(db_path=str(tmp_path / "test_dup.sqlite3"))
        service = KnowledgeService(storage_root=tmp_path, persistence=store)
        docx_bytes = make_test_docx("Duplicate Policy Test Content")

        doc1 = service.ingest_document(docx_bytes, "policy_v1.docx", "app_dup_test")
        assert doc1.status == DocumentStatus.READY

        doc2 = service.ingest_document(docx_bytes, "policy_v1_renamed.docx", "app_dup_test")
        assert doc2.status == DocumentStatus.DUPLICATE
        assert doc2.content_hash == doc1.content_hash

    def test_document_deletion_removes_chunks_and_files(self, tmp_path):
        """Deleting a document safely removes its records, chunks, and disk storage file."""
        store = KnowledgePersistenceStore(db_path=str(tmp_path / "test_del.sqlite3"))
        service = KnowledgeService(storage_root=tmp_path, persistence=store)
        pdf_bytes = make_test_pdf("Temporary Document to be Deleted.")

        doc = service.ingest_document(pdf_bytes, "temp_doc.pdf", "app_del_test")
        assert doc.status == DocumentStatus.READY
        assert Path(doc.storage_path).exists()

        # Delete
        success = service.delete_document(doc.document_id)
        assert success is True
        assert service.get_document(doc.document_id) is None
        assert len(service.list_chunks(doc.document_id)) == 0
        assert not Path(doc.storage_path).exists()

    def test_rest_api_knowledge_endpoints(self):
        """FastAPI REST endpoints for upload, listing, retrieval, and deletion function correctly."""
        headers = get_admin_headers()
        client = TestClient(app)

        pdf_bytes = make_test_pdf("API Enterprise Knowledge. Compliance requires annual certification renewal.")

        # 1. Upload Document via API
        files = {"file": ("api_policy.pdf", pdf_bytes, "application/pdf")}
        data = {"application_id": "app_api_rag", "chunk_size": 300, "chunk_overlap": 50}
        resp = client.post("/api/v2/knowledge/documents/upload", files=files, data=data, headers=headers)
        assert resp.status_code == 200, resp.text
        doc_data = resp.json()
        assert doc_data["filename"] == "api_policy.pdf"
        assert doc_data["status"] == "READY"
        doc_id = doc_data["document_id"]

        # 2. List Documents via API
        r_list = client.get("/api/v2/knowledge/documents?application_id=app_api_rag", headers=headers)
        assert r_list.status_code == 200
        docs = r_list.json()
        assert any(d["document_id"] == doc_id for d in docs)

        # 3. Get Document Chunks via API
        r_chunks = client.get(f"/api/v2/knowledge/documents/{doc_id}/chunks", headers=headers)
        assert r_chunks.status_code == 200
        chunks = r_chunks.json()
        assert len(chunks) >= 1

        # 4. Scoped Knowledge Retrieval via API
        ret_payload = {
            "application_id": "app_api_rag",
            "query": "annual certification renewal",
            "top_k": 3,
            "min_relevance_score": 0.01,
        }
        r_ret = client.post("/api/v2/knowledge/retrieve", json=ret_payload, headers=headers)
        assert r_ret.status_code == 200
        ret_data = r_ret.json()
        assert ret_data["status"] == "SUCCESS_WITH_RESULTS"
        assert len(ret_data["retrieved_chunks"]) >= 1
        assert ret_data["retrieved_chunks"][0]["filename"] == "api_policy.pdf"

        # 5. Delete Document via API
        r_del = client.delete(f"/api/v2/knowledge/documents/{doc_id}", headers=headers)
        assert r_del.status_code == 200
        assert r_del.json()["deleted"] is True

    def test_multi_tenant_authorization_isolation(self, tmp_path):
        """Users can only access documents in applications they have project memberships for."""
        import uuid
        from auth import UserRole, auth_repository

        client = TestClient(app)
        admin_token = create_access_token({"sub": "admin", "is_admin": True})
        admin_headers = {"Authorization": f"Bearer {admin_token}"}

        # Create User A for app_tenant_alpha
        user_a = auth_repository.create_user(
            user_id=f"usr_a_{uuid.uuid4().hex[:6]}",
            username=f"usera_{uuid.uuid4().hex[:6]}",
            email="usera@example.com",
            password="UserAPassword2026!",
            full_name="User Alpha",
        )
        auth_repository.assign_user_to_project(user_a.user_id, "app_tenant_alpha", UserRole.CERTIFIER)
        token_a = create_access_token(user=user_a)
        headers_a = {"Authorization": f"Bearer {token_a}"}

        # Create User B for app_tenant_beta
        user_b = auth_repository.create_user(
            user_id=f"usr_b_{uuid.uuid4().hex[:6]}",
            username=f"userb_{uuid.uuid4().hex[:6]}",
            email="userb@example.com",
            password="UserBPassword2026!",
            full_name="User Beta",
        )
        auth_repository.assign_user_to_project(user_b.user_id, "app_tenant_beta", UserRole.CERTIFIER)
        token_b = create_access_token(user=user_b)
        headers_b = {"Authorization": f"Bearer {token_b}"}

        pdf_bytes = make_test_pdf("Confidential Alpha Guidelines. Internal secret ALPHA_TOKEN_12345.")

        # 1. User A uploads to app_tenant_alpha -> ALLOW (200)
        files = {"file": ("alpha_doc.pdf", pdf_bytes, "application/pdf")}
        data_a = {"application_id": "app_tenant_alpha"}
        r_up_a = client.post("/api/v2/knowledge/documents/upload", files=files, data=data_a, headers=headers_a)
        assert r_up_a.status_code == 200
        doc_a_id = r_up_a.json()["document_id"]

        # 2. User B tries to upload to app_tenant_alpha -> DENY (403)
        files_b = {"file": ("malicious_doc.pdf", pdf_bytes, "application/pdf")}
        r_up_b_denied = client.post("/api/v2/knowledge/documents/upload", files=files_b, data=data_a, headers=headers_b)
        assert r_up_b_denied.status_code == 403, f"Expected 403, got {r_up_b_denied.status_code}: {r_up_b_denied.text}"

        # 3. User B tries to retrieve from app_tenant_alpha -> DENY (403)
        ret_b_payload = {"application_id": "app_tenant_alpha", "query": "ALPHA_TOKEN_12345"}
        r_ret_b_denied = client.post("/api/v2/knowledge/retrieve", json=ret_b_payload, headers=headers_b)
        assert r_ret_b_denied.status_code == 403, f"Expected 403, got {r_ret_b_denied.status_code}: {r_ret_b_denied.text}"

        # 4. User B tries to view or delete doc_a_id -> DENY (403)
        r_get_b_denied = client.get(f"/api/v2/knowledge/documents/{doc_a_id}", headers=headers_b)
        assert r_get_b_denied.status_code == 403
        r_del_b_denied = client.delete(f"/api/v2/knowledge/documents/{doc_a_id}", headers=headers_b)
        assert r_del_b_denied.status_code == 403

        # 5. User A can retrieve from app_tenant_alpha -> ALLOW (200)
        ret_a_payload = {"application_id": "app_tenant_alpha", "query": "ALPHA_TOKEN_12345"}
        r_ret_a = client.post("/api/v2/knowledge/retrieve", json=ret_a_payload, headers=headers_a)
        assert r_ret_a.status_code == 200
        assert len(r_ret_a.json()["retrieved_chunks"]) >= 1

        # Clean up
        client.delete(f"/api/v2/knowledge/documents/{doc_a_id}", headers=admin_headers)

    def test_txt_and_text_upload_rejected_by_api(self):
        """API strictly rejects .txt and .text uploads with HTTP 400."""
        headers = get_admin_headers()
        client = TestClient(app)

        files_txt = {"file": ("policy.txt", b"Some plain text rules.", "text/plain")}
        r_txt = client.post("/api/v2/knowledge/documents/upload", files=files_txt, data={"application_id": "app_txt"}, headers=headers)
        assert r_txt.status_code == 400
        assert "Only PDF and DOCX documents are supported" in r_txt.json()["detail"]

        files_text = {"file": ("policy.text", b"Some plain text rules.", "text/plain")}
        r_text = client.post("/api/v2/knowledge/documents/upload", files=files_text, data={"application_id": "app_txt"}, headers=headers)
        assert r_text.status_code == 400
        assert "Only PDF and DOCX documents are supported" in r_text.json()["detail"]

    def test_safe_error_masking_no_exception_leakage(self, monkeypatch):
        """Internal server exceptions do not leak stack traces, database details, or filesystem paths."""
        headers = get_admin_headers()
        client = TestClient(app)

        # Mock an unexpected crash in ingestion
        def mock_ingest(*args, **kwargs):
            raise RuntimeError("sqlite3.OperationalError: database disk image is malformed at /var/secrets/db.sqlite")

        monkeypatch.setattr(knowledge_service, "ingest_document", mock_ingest)

        pdf_bytes = make_test_pdf("Standard text content.")
        files = {"file": ("doc.pdf", pdf_bytes, "application/pdf")}
        resp = client.post("/api/v2/knowledge/documents/upload", files=files, data={"application_id": "app_err"}, headers=headers)
        assert resp.status_code == 500
        detail = resp.json()["detail"]
        # Must be generic and safe
        assert detail == "Document ingestion failed."
        assert "sqlite3" not in detail
        assert "malformed" not in detail
        assert "/var/secrets" not in detail

    def test_retrieval_fingerprint_policy_differentiation(self, tmp_path):
        """Retrieval fingerprints are deterministic and differ when retrieval policy changes."""
        from knowledge.models import KnowledgeQuery, RetrievalPolicy
        store = KnowledgePersistenceStore(db_path=str(tmp_path / "test_fp.sqlite3"))
        service = KnowledgeService(storage_root=tmp_path, persistence=store)

        pdf_bytes = make_test_pdf("Corporate Refund Policy. Section 1. All refunds processed within 5 business days.")
        service.ingest_document(pdf_bytes, "refund.pdf", "app_fp_test")

        # 1. Retrieval with top_k = 3
        ctx1 = service.retrieve(
            query="refunds processed",
            application_id="app_fp_test",
            top_k=3,
            policy=RetrievalPolicy(top_k=3, min_relevance_score=0.01),
        )

        # 2. Identical retrieval with top_k = 3 -> same fingerprint
        ctx2 = service.retrieve(
            query="refunds processed",
            application_id="app_fp_test",
            top_k=3,
            policy=RetrievalPolicy(top_k=3, min_relevance_score=0.01),
        )
        assert ctx1.retrieval_fingerprint == ctx2.retrieval_fingerprint
        assert len(ctx1.retrieval_fingerprint) == 64

        # 3. Retrieval with top_k = 10 -> different fingerprint
        ctx3 = service.retrieve(
            query="refunds processed",
            application_id="app_fp_test",
            top_k=10,
            policy=RetrievalPolicy(top_k=10, min_relevance_score=0.01),
        )
        assert ctx1.retrieval_fingerprint != ctx3.retrieval_fingerprint

    def test_post_delete_retrieval_consistency(self, tmp_path):
        """After a document is deleted, retrieval against its scope returns SUCCESS_NO_RESULTS."""
        store = KnowledgePersistenceStore(db_path=str(tmp_path / "test_post_del.sqlite3"))
        service = KnowledgeService(storage_root=tmp_path, persistence=store)

        pdf_bytes = make_test_pdf("Exclusive Project Phoenix Blueprint and Architecture.")
        doc = service.ingest_document(pdf_bytes, "phoenix.pdf", "app_phoenix")
        assert doc.status == DocumentStatus.READY

        # Prior to deletion -> found
        ctx_before = service.retrieve(query="Phoenix Blueprint", application_id="app_phoenix")
        assert ctx_before.status == RetrievalStatus.SUCCESS_WITH_RESULTS
        assert len(ctx_before.retrieved_chunks) >= 1

        # Delete document
        deleted = service.delete_document(doc.document_id)
        assert deleted is True

        # After deletion -> no results
        ctx_after = service.retrieve(query="Phoenix Blueprint", application_id="app_phoenix")
        assert ctx_after.status == RetrievalStatus.SUCCESS_NO_RESULTS
        assert len(ctx_after.retrieved_chunks) == 0

